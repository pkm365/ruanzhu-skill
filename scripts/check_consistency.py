#!/usr/bin/env python3
"""
提交前自检：三份材料的名称/版本/权利人一致性 + 各项硬性指标。

用法:
  uv run --with python-docx --with pyyaml python3 check_consistency.py 正式资料/ [--yaml 草稿/申请表信息.yaml]

检查项:
  1. 申请表 / 手册封面+页眉 / 代码页眉 中的 全称+版本号 完全一致
  2. 手册封面著作权人 == 申请表权利人名称
  3. 申请表: 主要功能 500–1300 字，技术特点 ≤100 字，必填项非空
  4. 手册: 无【截图预留】/【截图缺失】残留；一级标题数、插图数；正文段落估算页数 ≥20
  5. 手册一级标题（模块名）是否在申请表"主要功能"中被提到
  6. 代码: 每页 50 行、总页数 ≤60、末页标注；代码文档索引中的语言 ⊆ 申请表"编程语言"
"""
import argparse, re, sys
from pathlib import Path
from docx import Document
import yaml

EXT_LANG = {".java": "Java", ".py": "Python", ".js": "JavaScript", ".jsx": "JavaScript", ".ts": "TypeScript", ".tsx": "TypeScript",
            ".vue": "Vue", ".html": "HTML", ".css": "CSS", ".scss": "CSS", ".less": "CSS", ".go": "Go", ".cs": "C#", ".php": "PHP", ".kt": "Kotlin", ".sql": "SQL"}
VUE_IMPLIES = {"JavaScript", "HTML", "CSS"}

def header_text(doc):
    return "".join(p.text for p in doc.sections[0].header.paragraphs).split("\t")[0].strip()

def cjk_len(s): return len(re.sub(r"\s", "", s or ""))

def main():
    ap = argparse.ArgumentParser(); ap.add_argument("dir"); ap.add_argument("--yaml", default=None)
    a = ap.parse_args(); d = Path(a.dir)
    problems, infos = [], []
    P = lambda m: problems.append("❌ " + m)
    I = lambda m: infos.append("ℹ️  " + m)

    form = next(d.glob("*-申请表.docx"), None)
    manual = next(d.glob("*_操作手册.docx"), None)
    code = next(d.glob("*_软件代码.docx"), None)
    for label, f in (("申请表", form), ("操作手册", manual), ("软件代码", code)):
        if f is None: P(f"缺少 {label} docx")
    data = {}
    if a.yaml: data = yaml.safe_load(Path(a.yaml).read_text(encoding="utf-8")) or {}
    elif form:
        doc = Document(form)
        for t in doc.tables:
            for row in t.rows:
                c = row.cells
                if len(c) >= 2:
                    k = re.split(r"[（(\n]", c[0].text.strip())[0].strip(); v = c[-1].text.strip()
                    if k and k not in data: data[k] = v
    name, ver, owner = data.get("软件全称", ""), data.get("版本号", ""), data.get("权利人名称", "")
    full = f"{name}{ver}"
    I(f"申请表: 全称「{name}」 版本「{ver}」 权利人「{owner}」")

    # 3 申请表指标
    n = cjk_len(data.get("软件的主要功能", ""))
    if not 500 <= n <= 1300: P(f"主要功能 {n} 字，需 500–1300")
    else: I(f"主要功能 {n} 字 ✓")
    t = cjk_len(data.get("软件的技术特点", ""))
    if t > 100: P(f"技术特点 {t} 字，需 ≤100")
    if ver and not re.fullmatch(r"V\d+(\.\d+)*", ver): P(f"版本号「{ver}」应为 V1.0 形式")

    # 1/2 手册
    if manual:
        m = Document(manual); h = header_text(m)
        if h != full: P(f"手册页眉「{h}」≠ 申请表「{full}」")
        paras = [p.text.strip() for p in m.paragraphs]
        cover_owner = next((p.replace("著作权人：", "") for p in paras if p.startswith("著作权人：")), "")
        if cover_owner != owner: P(f"手册封面著作权人「{cover_owner}」≠ 申请表权利人「{owner}」")
        cover_ver = next((p.replace("版本：", "") for p in paras if p.startswith("版本：")), "")
        if cover_ver != ver: P(f"手册封面版本「{cover_ver}」≠「{ver}」")
        if paras and paras[0] == "" and name not in paras[:12]: P("手册封面未见软件全称")
        leftovers = [p for p in paras if p.startswith(("【截图预留", "【截图缺失"))]
        if leftovers: P(f"手册仍有 {len(leftovers)} 处截图占位: " + " / ".join(x[:30] for x in leftovers[:5]))
        h1 = [p.text.strip() for p in m.paragraphs if p.style.name == "Heading 1"]
        imgs = m.element.body.xml.count("<pic:pic")
        body_chars = sum(len(p) for p in paras)
        pdf = manual.with_suffix(".pdf")
        if pdf.exists():
            try:
                from pypdf import PdfReader
                real = len(PdfReader(str(pdf)).pages) - 2  # 去封面、目录
                I(f"手册: 一级标题 {len(h1)} 个，插图 {imgs} 张，正文约 {body_chars} 字，PDF 实际 {real} 页（不含封面目录）")
                if real < 20: P(f"手册 PDF 仅 {real} 页，要求除目录外 ≥20 页")
            except ImportError:
                I("装 pypdf 可按 PDF 真实页数校验：uv run --with pypdf ...")
        else:
            est_pages = body_chars / 900 + imgs * 0.45
            I(f"手册: 一级标题 {len(h1)} 个，插图 {imgs} 张，正文约 {body_chars} 字，估算 {est_pages:.0f} 页（含图，不含封面目录；转 PDF 后按真实页数校验）")
            if est_pages < 20: P(f"手册估算仅 {est_pages:.0f} 页，要求除目录外 ≥20 页")
        if imgs == 0: P("手册没有任何截图")
        func = data.get("软件的主要功能", "")
        for title in h1:
            core = re.sub(r"^[\d.、\s一二三四五六七八九十]+", "", title)
            core = re.sub(r"(系统|模块|管理|功能|设置)$", "", core) or core
            if core and core.replace(" ", "") not in func.replace(" ", ""): I(f"手册章节「{title}」在申请表主要功能里未出现，确认是否需要对应")
        for bad in ("import ", "public class", "function(", "SELECT ", "def "):
            if any(bad in p for p in paras): P(f"手册疑似含代码片段「{bad.strip()}」，说明书禁止出现功能函数代码"); break

    # 6 代码
    if code:
        c = Document(code); h = header_text(c)
        if h != full: P(f"代码页眉「{h}」≠ 申请表「{full}」")
        pages, cur = [], []
        for p in c.paragraphs:
            if p._element.pPr is not None and p._element.pPr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pageBreakBefore") is not None and cur:
                pages.append(cur); cur = []
            if p.text.strip(): cur.append(p.text)
        if cur: pages.append(cur)
        last_note = any("本页为源代码最后一页" in x for x in pages[-1]) if pages else False
        code_pages = [pg if i < len(pages) - 1 else [x for x in pg if "本页为源代码最后一页" not in x] for i, pg in enumerate(pages)]
        bad = [i + 1 for i, pg in enumerate(code_pages[:-1]) if len(pg) != 50]
        I(f"代码: {len(pages)} 页，末页 {len(code_pages[-1]) if pages else 0} 行" + ("，已标注末页" if last_note else ""))
        if bad: P(f"代码非末页行数≠50: 第 {bad[:10]} 页")
        if len(pages) > 60: P(f"代码 {len(pages)} 页 > 60")
        if pages and len(code_pages[-1]) < 50 and not last_note: P("末页不足 50 行但未标注「本页为源代码最后一页」")
        idx = d / "代码文档索引.md"
        if idx.exists():
            exts = set(re.findall(r"`[^`]+?(\.[a-zA-Z]+)`", idx.read_text(encoding="utf-8")))
            langs = set()
            for e in exts:
                l = EXT_LANG.get(e.lower())
                if l == "Vue": langs |= VUE_IMPLIES
                elif l: langs.add(l)
            declared = data.get("编程语言", "")
            missing = [l for l in langs if l not in declared]
            if missing: P(f"代码涉及 {sorted(langs)}，申请表编程语言「{declared}」缺少 {missing}")
            else: I(f"编程语言 {sorted(langs)} ⊆ 申请表 ✓")

    print("\n".join(infos)); print()
    if problems: print("\n".join(problems)); print(f"\n共 {len(problems)} 个问题，修正后重跑。"); sys.exit(1)
    print("✅ 自检通过。仍需人工确认：截图完整含标题栏、权属/完成日期真实、Word 目录已更新。")

if __name__ == "__main__":
    main()
