"""python-docx 公共工具：页眉页码字段、字体、分页、页面设置。"""
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm

def set_font(run, name="宋体", size=None, bold=None, ascii_name=None):
    run.font.name = ascii_name or name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts"); rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:ascii"), ascii_name or name)
    rfonts.set(qn("w:hAnsi"), ascii_name or name)
    if size: run.font.size = Pt(size)
    if bold is not None: run.font.bold = bold

def set_style_font(style, name="宋体", size=None, ascii_name=None):
    style.font.name = ascii_name or name
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts"); rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:ascii"), ascii_name or name)
    rfonts.set(qn("w:hAnsi"), ascii_name or name)
    if size: style.font.size = Pt(size)

def add_field(paragraph, instr: str, placeholder="1"):
    """插入 Word 域（PAGE / NUMPAGES / TOC ...）"""
    r1 = paragraph.add_run(); fc1 = OxmlElement("w:fldChar"); fc1.set(qn("w:fldCharType"), "begin"); r1._element.append(fc1)
    r2 = paragraph.add_run(); it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = f" {instr} "; r2._element.append(it)
    r3 = paragraph.add_run(); fc2 = OxmlElement("w:fldChar"); fc2.set(qn("w:fldCharType"), "separate"); r3._element.append(fc2)
    r4 = paragraph.add_run(placeholder)
    r5 = paragraph.add_run(); fc3 = OxmlElement("w:fldChar"); fc3.set(qn("w:fldCharType"), "end"); r5._element.append(fc3)
    return [r1, r2, r3, r4, r5]

def setup_header(section, left_text: str, font_size=9, font="宋体"):
    """页眉：左侧 软件全称+版本号，右侧 第 N 页（右上角页码）"""
    header = section.header
    header.is_linked_to_previous = False
    p = header.paragraphs[0]
    for r in list(p.runs): r._element.getparent().remove(r._element)
    text_width = section.page_width - section.left_margin - section.right_margin
    p.paragraph_format.tab_stops.add_tab_stop(text_width, WD_TAB_ALIGNMENT.RIGHT)
    r = p.add_run(left_text); set_font(r, font, font_size)
    r = p.add_run("\t第 "); set_font(r, font, font_size)
    for r in add_field(p, "PAGE"): set_font(r, font, font_size)
    r = p.add_run(" 页"); set_font(r, font, font_size)
    # 清空页脚
    for fp in section.footer.paragraphs:
        for r in list(fp.runs): r._element.getparent().remove(r._element)

def setup_page(section, top, bottom, left, right, header_dist=0.8, footer_dist=0.6):
    section.page_width = Cm(21.0); section.page_height = Cm(29.7)
    section.top_margin = Cm(top); section.bottom_margin = Cm(bottom)
    section.left_margin = Cm(left); section.right_margin = Cm(right)
    section.header_distance = Cm(header_dist); section.footer_distance = Cm(footer_dist)

def remove_doc_grid(section):
    """去掉默认模板的 18pt 文档网格，否则行距被强制拉大，50 行装不下一页"""
    sectPr = section._sectPr
    for g in sectPr.findall(qn("w:docGrid")):
        sectPr.remove(g)

def page_break_before(paragraph):
    pPr = paragraph._element.get_or_add_pPr()
    el = OxmlElement("w:pageBreakBefore"); pPr.append(el)

def enable_update_fields(doc):
    """打开文档时提示更新域（用于目录）"""
    settings = doc.settings.element
    uf = settings.find(qn("w:updateFields"))
    if uf is None:
        uf = OxmlElement("w:updateFields"); settings.append(uf)
    uf.set(qn("w:val"), "true")
